"""Integration test suite for font embedding feature."""

from __future__ import annotations

import os
import struct
import zipfile
from io import BytesIO
from xml.etree import ElementTree

import pytest

from pptx import Presentation

NSMAP = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}

TEST_FILES_DIR = os.path.join(os.path.dirname(__file__), "test_files")
CALIBRIZ_PATH = os.path.join(TEST_FILES_DIR, "calibriz.ttf")


def _read_font_bytes():
    with open(CALIBRIZ_PATH, "rb") as f:
        return f.read()


class DescribeFontEmbeddingIntegration:
    """End-to-end tests for embedding fonts into a .pptx package."""

    def it_embeds_a_font_into_the_pptx_package(self):
        prs = Presentation()
        font_bytes = _read_font_bytes()

        prs.embed_font("Calibri", font_bytes)

        stream = BytesIO()
        prs.save(stream)
        stream.seek(0)

        with zipfile.ZipFile(stream) as z:
            # font part exists in the ZIP
            font_names = [n for n in z.namelist() if n.startswith("ppt/fonts/")]
            assert len(font_names) == 1
            assert font_names[0] == "ppt/fonts/font1.fntdata"

            # font data is in EOT format (not raw TTF)
            fntdata = z.read("ppt/fonts/font1.fntdata")
            eot_size, font_data_size, version = struct.unpack_from("<III", fntdata, 0)
            assert version == 0x00020001  # EOT version
            magic = struct.unpack_from("<H", fntdata, 34)[0]
            assert magic == 0x504C  # EOT magic number
            assert eot_size == len(fntdata)
            assert font_data_size == len(font_bytes)
            # raw TTF is appended after the header
            assert fntdata[-len(font_bytes):] == font_bytes

            # content types includes fntdata
            ct_xml = z.read("[Content_Types].xml").decode("utf-8")
            assert "application/x-fontdata" in ct_xml

            # presentation.xml has embeddedFontLst
            prs_xml = z.read("ppt/presentation.xml")
            root = ElementTree.fromstring(prs_xml)
            font_lst = root.find("p:embeddedFontLst", NSMAP)
            assert font_lst is not None

            embedded_fonts = font_lst.findall("p:embeddedFont", NSMAP)
            assert len(embedded_fonts) == 1

            font_elm = embedded_fonts[0].find("p:font", NSMAP)
            assert font_elm is not None
            assert font_elm.get("typeface") == "Calibri"
            assert font_elm.get("pitchFamily") == "2"
            assert font_elm.get("charset") == "0"

            regular_elm = embedded_fonts[0].find("p:regular", NSMAP)
            assert regular_elm is not None
            r_id = regular_elm.get("{%s}id" % NSMAP["r"])
            assert r_id is not None

            # verify the relationship exists in the rels file
            rels_xml = z.read("ppt/_rels/presentation.xml.rels").decode("utf-8")
            assert r_id in rels_xml
            assert "relationships/font" in rels_xml

    def it_embeds_multiple_variants_of_the_same_font(self):
        prs = Presentation()
        font_bytes = _read_font_bytes()

        prs.embed_font("Calibri", font_bytes, bold=False, italic=False)
        prs.embed_font("Calibri", font_bytes, bold=True, italic=False)

        stream = BytesIO()
        prs.save(stream)
        stream.seek(0)

        with zipfile.ZipFile(stream) as z:
            # two font parts
            font_names = [n for n in z.namelist() if n.startswith("ppt/fonts/")]
            assert len(font_names) == 2

            # single embeddedFont entry with both regular and bold
            prs_xml = z.read("ppt/presentation.xml")
            root = ElementTree.fromstring(prs_xml)
            font_lst = root.find("p:embeddedFontLst", NSMAP)
            embedded_fonts = font_lst.findall("p:embeddedFont", NSMAP)
            assert len(embedded_fonts) == 1

            assert embedded_fonts[0].find("p:regular", NSMAP) is not None
            assert embedded_fonts[0].find("p:bold", NSMAP) is not None

    def it_embeds_multiple_different_fonts(self):
        prs = Presentation()
        font_bytes = _read_font_bytes()

        prs.embed_font("FontA", font_bytes)
        prs.embed_font("FontB", font_bytes)

        stream = BytesIO()
        prs.save(stream)
        stream.seek(0)

        with zipfile.ZipFile(stream) as z:
            prs_xml = z.read("ppt/presentation.xml")
            root = ElementTree.fromstring(prs_xml)
            font_lst = root.find("p:embeddedFontLst", NSMAP)
            embedded_fonts = font_lst.findall("p:embeddedFont", NSMAP)
            assert len(embedded_fonts) == 2

            typefaces = [
                ef.find("p:font", NSMAP).get("typeface") for ef in embedded_fonts
            ]
            assert typefaces == ["FontA", "FontB"]

    def it_produces_a_valid_pptx_that_can_be_reopened(self):
        prs = Presentation()
        font_bytes = _read_font_bytes()
        prs.embed_font("Calibri", font_bytes)

        stream = BytesIO()
        prs.save(stream)
        stream.seek(0)

        # re-open should not raise
        prs2 = Presentation(stream)
        assert prs2.slides is not None
