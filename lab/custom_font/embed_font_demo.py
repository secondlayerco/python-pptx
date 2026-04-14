"""Demo script for font embedding.

Generates a .pptx with an embedded custom font. Open the output file on a
machine that does NOT have the font installed to verify it renders correctly.

Usage:
    python lab/embed_font_demo.py path/to/font.ttf "Font Family Name"

Example with the test font already in the repo:
    python lab/embed_font_demo.py tests/test_files/calibriz.ttf "Calibri"

To test with a Google Fonts download (e.g. Beth Ellen):
    python lab/embed_font_demo.py BethEllen-Regular.ttf "Beth Ellen"
"""

from __future__ import annotations

import sys

from pptx import Presentation
from pptx.util import Inches, Pt


def main(font_path: str, typeface: str) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout

    # add a text box with the custom font
    txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    tf = txBox.text_frame
    run = tf.paragraphs[0].add_run()
    run.text = f"Hello from {typeface}!"
    run.font.name = typeface
    run.font.size = Pt(36)

    # add a second line with default font for comparison
    run2 = tf.add_paragraph().add_run()
    run2.text = "This line uses the default theme font."
    run2.font.size = Pt(18)

    # embed the font
    with open(font_path, "rb") as f:
        font_bytes = f.read()
    prs.embed_font(typeface, font_bytes)

    out = "test_custom_font.pptx"
    prs.save(out)
    print(f"Saved {out} with embedded font '{typeface}' from {font_path}")


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python lab/embed_font_demo.py <font.ttf> <typeface name>")
        sys.exit(1)
    main(sys.argv[1], sys.argv[2])
