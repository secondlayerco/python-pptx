"""Debug script: generate multiple PPTX variants to isolate font issues."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.parts.font import FontPart
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.eot import ttf_to_eot
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement

FONT_PATH = "lab/custom_font/BitcountGridDouble.ttf"

with open(FONT_PATH, "rb") as f:
    font_bytes = f.read()


def make_prs():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    run = txBox.text_frame.paragraphs[0].add_run()
    run.text = "Hello custom font"
    run.font.name = "Bitcount Grid Double"
    run.font.size = Pt(36)
    return prs


# Test 5: EOT + embedTrueTypeFonts=1
def test_eot_with_flag():
    prs = make_prs()
    prs.part._element.set("embedTrueTypeFonts", "1")
    prs.embed_font("Bitcount Grid Double", font_bytes)
    prs.save("test_5_eot_flag.pptx")
    print("Saved test_5_eot_flag.pptx (EOT + embedTrueTypeFonts=1)")


# Test 6: Raw TTF + embedTrueTypeFonts=1
def test_raw_with_flag():
    prs = make_prs()
    prs.part._element.set("embedTrueTypeFonts", "1")
    part = prs.part
    font_part = FontPart.new(font_bytes, part.package)
    rId = part.relate_to(font_part, RT.FONT)
    lst = part._element.get_or_add_embeddedFontLst()
    ef = OxmlElement("p:embeddedFont")
    font_elm = OxmlElement("p:font")
    font_elm.set("typeface", "Bitcount Grid Double")
    font_elm.set("pitchFamily", "2")
    font_elm.set("charset", "0")
    ef.append(font_elm)
    style_elm = OxmlElement("p:regular")
    style_elm.set(qn("r:id"), rId)
    ef.append(style_elm)
    lst.append(ef)
    prs.save("test_6_raw_flag.pptx")
    print("Saved test_6_raw_flag.pptx (raw TTF + embedTrueTypeFonts=1)")


test_eot_with_flag()
test_raw_with_flag()
print("\nTry both in PowerPoint.")
