"""Generate test PPTX files for debugging font embedding."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.parts.font import FontPart
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement

FONT_PATH = "lab/custom_font/BitcountGridDouble.ttf"
TYPEFACE = "Bitcount Grid Double"


def gen_no_font():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    run = txBox.text_frame.paragraphs[0].add_run()
    run.text = "No embedded font test"
    run.font.size = Pt(36)
    prs.save("test_no_font.pptx")
    print("Saved test_no_font.pptx (no font embedding)")


def gen_raw_ttf():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    run = txBox.text_frame.paragraphs[0].add_run()
    run.text = "Hello from raw TTF embed"
    run.font.name = TYPEFACE
    run.font.size = Pt(36)

    with open(FONT_PATH, "rb") as f:
        font_bytes = f.read()

    part = prs.part
    font_part = FontPart.new(font_bytes, part.package)
    rId = part.relate_to(font_part, RT.FONT)
    embeddedFontLst = part._element.get_or_add_embeddedFontLst()
    ef = OxmlElement("p:embeddedFont")
    font_elm = OxmlElement("p:font")
    font_elm.set("typeface", TYPEFACE)
    font_elm.set("pitchFamily", "2")
    font_elm.set("charset", "0")
    ef.append(font_elm)
    style_elm = OxmlElement("p:regular")
    style_elm.set(qn("r:id"), rId)
    ef.append(style_elm)
    embeddedFontLst.append(ef)
    prs.save("test_raw_ttf.pptx")
    print("Saved test_raw_ttf.pptx (raw TTF, no EOT)")


def gen_eot():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    run = txBox.text_frame.paragraphs[0].add_run()
    run.text = "Hello from EOT embed"
    run.font.name = TYPEFACE
    run.font.size = Pt(36)

    with open(FONT_PATH, "rb") as f:
        font_bytes = f.read()
    prs.embed_font(TYPEFACE, font_bytes)
    prs.save("test_eot.pptx")
    print("Saved test_eot.pptx (EOT conversion)")


if __name__ == "__main__":
    gen_no_font()
    gen_raw_ttf()
    gen_eot()
